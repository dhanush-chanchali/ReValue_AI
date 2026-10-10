import sys, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Palette
    C_BG = RGBColor(8, 13, 26)       # #080d1a
    C_CARD = RGBColor(13, 20, 36)    # #0d1424
    C_GOLD = RGBColor(212, 175, 55)   # #d4af37
    C_EMERALD = RGBColor(0, 229, 153) # #00e599
    C_WHITE = RGBColor(241, 245, 249) # #f1f5f9
    C_MUTED = RGBColor(148, 163, 184) # #94a3b8
    C_DARK_CARD = RGBColor(18, 27, 48)

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text):
        # Category Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11), Inches(0.4))
        tf = tag_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = tag_text.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_EMERALD

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.8))
        tf2 = title_box.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(24)
        p2.font.bold = True
        p2.font.color.rgb = C_WHITE

    # ─────────────────────────────────────────────────────────────
    # SLIDE 1: Title Slide
    # ─────────────────────────────────────────────────────────────
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Accent banner box
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(5.8), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = C_DARK_CARD
    badge.line.color.rgb = C_GOLD
    p = badge.text_frame.paragraphs[0]
    p.text = "⚡ PRODUCTION MACHINE LEARNING PLATFORM  •  CO-1 TO CO-6"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = C_GOLD
    p.alignment = PP_ALIGN.CENTER

    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(2.2))
    tf = t_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "ReValue AI"
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    p2 = tf.add_paragraph()
    p2.text = "Intelligent Electronics Resale Valuation & MLOps Platform"
    p2.font.size = Pt(22)
    p2.font.color.rgb = C_EMERALD

    # Subtitle / Overview Card
    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), Inches(11.7), Inches(2.2))
    card.fill.solid()
    card.fill.fore_color.rgb = C_CARD
    card.line.color.rgb = RGBColor(40, 50, 75)
    ctf = card.text_frame
    ctf.word_wrap = True
    cp = ctf.paragraphs[0]
    cp.text = "Key Project Highlights & Architectural Pillars:"
    cp.font.size = Pt(13)
    cp.font.bold = True
    cp.font.color.rgb = C_GOLD

    bullets = [
        "🎯 Production Champion Regressor: HistGradientBoosting (R² = 0.9877, Test MAE = ₹2,305 / $27.44) with 256-bin split acceleration.",
        "🛡️ Unsupervised Fraud Guardrail: Isolation Forest anomaly screening catches fraudulent trade-in submissions before inference.",
        "🌐 Dual-Engine Deployment: High-throughput FastAPI REST microservice + Streamlit Cloud interactive valuation studio.",
        "📈 Continuous MLOps Telemetry: Live SQLite logging, Two-Sample Kolmogorov-Smirnov (KS) drift testing, automated retraining loop."
    ]
    for b in bullets:
        bp = ctf.add_paragraph()
        bp.text = b
        bp.font.size = Pt(11)
        bp.font.color.rgb = C_WHITE
        bp.space_before = Pt(4)

    s1.notes_slide.notes_text_frame.text = (
        "[TIMING: 0:00 - 1:00]\n"
        "TALKING SCRIPT:\n"
        "'Good morning respected professors and evaluators. Today I am proud to present ReValue AI — an end-to-end, "
        "production-grade Machine Learning platform designed to eliminate information asymmetry in the secondhand electronics marketplace.\n"
        "ReValue AI is not a simple notebook prototype. It is a full lifecycle engineering system that rigorously benchmarks "
        "regularized linear baselines against advanced ensemble trees, integrates unsupervised K-Means clustering and Isolation Forest fraud auditing, "
        "and runs live in production with continuous Kolmogorov-Smirnov drift telemetry, covering Course Outcomes CO-1 through CO-6.'"
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 2: Problem Statement & Motivation
    # ─────────────────────────────────────────────────────────────
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Market Opportunity & Engineering Need", "The Secondhand Electronics Dilemma")

    cols = [
        ("Information Asymmetry", "Buyers and sellers rely on subjective intuition or fragmented marketplace listings without an empirical baseline for true hardware residual worth.", C_GOLD),
        ("Discontinuous Hardware Degradation", "Device depreciation is strictly non-linear. Battery degradation sub-80%, physical screen cracks, and missing original chargers create sharp cliff drops.", C_EMERALD),
        ("Fraudulent & Bogus Quotes", "Resale portals suffer from quote manipulation (e.g. ₹10,000 for an iPhone 16 Pro Max) that pollutes pricing engines and creates seller fraud.", C_WHITE)
    ]
    for i, (col_title, col_desc, col_color) in enumerate(cols):
        c_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 4.0), Inches(1.8), Inches(3.7), Inches(4.8))
        c_box.fill.solid()
        c_box.fill.fore_color.rgb = C_CARD
        c_box.line.color.rgb = col_color
        ctf = c_box.text_frame
        ctf.word_wrap = True
        cp = ctf.paragraphs[0]
        cp.text = col_title
        cp.font.size = Pt(16)
        cp.font.bold = True
        cp.font.color.rgb = col_color
        
        cp2 = ctf.add_paragraph()
        cp2.text = col_desc
        cp2.font.size = Pt(12)
        cp2.font.color.rgb = C_WHITE
        cp2.space_before = Pt(14)

    s2.notes_slide.notes_text_frame.text = (
        "[TIMING: 1:00 - 2:00]\n"
        "TALKING SCRIPT:\n"
        "'The secondhand smartphone and laptop industry is a multi-billion dollar market plagued by three core challenges: "
        "First, information asymmetry. Neither buyers nor sellers know what a device is objectively worth once it leaves the showroom.\n"
        "Second, electronics depreciation is inherently non-linear. A laptop battery below 80% health drops trade-in value drastically, "
        "creating decision boundaries that simple linear formulas fail to capture.\n"
        "Third, resale portals are vulnerable to fraudulent or corrupted quotes. ReValue AI solves all three using a dual-engine architecture.'"
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 3: System Architecture & ML Lifecycle (CO-1)
    # ─────────────────────────────────────────────────────────────
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "CO-1: ML Lifecycle & System Architecture", "Zero Training-Serving Skew Microservice Architecture")

    stages = [
        ("1. Data Pipeline", "src/data_pipeline.py\nStrict 70/15/15 3-way partition (Train / Val / Test). Validation contracts with Pydantic."),
        ("2. Atomic Preprocessor", "src/features.py\nColumnTransformer bundling One-Hot, Ordinal wear grading, and isotropic StandardScaler into a single artifact."),
        ("3. Model Benchmarking", "src/models_training.py\n5-Fold cross-validation across 6 algorithms: OLS, Ridge, Lasso, ElasticNet, Random Forest, HistGBDT."),
        ("4. Serving Microservice", "src/api.py\nFastAPI asynchronous REST service with sub-10ms response times and 90% confidence interval generation."),
        ("5. Telemetry & Retraining", "src/monitoring.py & retrain.py\nSQLite inference database, 2-Sample Kolmogorov-Smirnov drift detection, automated retraining loop.")
    ]
    for i, (stg_title, stg_desc) in enumerate(stages):
        box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8 + i * 1.0), Inches(11.7), Inches(0.85))
        box.fill.solid()
        box.fill.fore_color.rgb = C_CARD
        box.line.color.rgb = C_GOLD if i == 1 else (C_EMERALD if i == 2 else RGBColor(40, 50, 75))
        btf = box.text_frame
        btf.word_wrap = True
        bp = btf.paragraphs[0]
        bp.text = stg_title + "  ➔  "
        bp.font.size = Pt(13)
        bp.font.bold = True
        bp.font.color.rgb = C_EMERALD
        
        run = bp.add_run()
        run.text = stg_desc
        run.font.size = Pt(11)
        run.font.bold = False
        run.font.color.rgb = C_WHITE

    s3.notes_slide.notes_text_frame.text = (
        "[TIMING: 2:00 - 3:15]\n"
        "TALKING SCRIPT:\n"
        "'Moving to Course Outcome CO-1: The ML Lifecycle. The single most fatal bug in production ML is Training-Serving Skew — "
        "where preprocessing logic in training drifts away from the serving API.\n"
        "In ReValue AI, we solved this completely. We encapsulate median imputation, ordinal cosmetic scoring, one-hot category encoders, "
        "and standard scalers inside an atomic Scikit-Learn ColumnTransformer pipeline exported as best_model_pipeline.joblib.\n"
        "Both our FastAPI service and our Streamlit dashboard ingest raw dictionaries directly into this pipeline. "
        "There is zero manual feature recalculation in production, guaranteeing mathematical zero skew.'"
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 4: Algorithms Directory: 11 Algorithms Mapped
    # ─────────────────────────────────────────────────────────────
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Comprehensive Algorithm Mapping", "11 Algorithms Across the ML Lifecycle")

    tbl_shape = s4.shapes.add_table(7, 4, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    table = tbl_shape.table
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(2.3)
    table.columns[2].width = Inches(2.2)
    table.columns[3].width = Inches(4.0)

    headers = ["Algorithm", "Learning Paradigm", "Code Location", "Why Used in ReValue AI"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_DARK_CARD
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_GOLD

    rows_data = [
        ("HistGradientBoosting (Champion)", "Supervised Regression (Boosting)", "src/models_training.py", "Production Champion Model (R²=0.9877, MAE=₹2,305). 256-binning."),
        ("Random Forest Regressor", "Supervised Regression (Bagging)", "src/models_training.py", "Ensemble baseline (150 trees, sqrt(p) feature decorrelation)."),
        ("Ridge Regression (L2)", "Supervised Linear (L2 Penalized)", "src/models_training.py", "Linear baseline; resolves ill-conditioned collinear hardware inversion."),
        ("Lasso Regression (L1)", "Supervised Linear (L1 Sparse)", "src/models_training.py", "Automated feature sparsity; drives redundant coefficients to zero."),
        ("K-Means Clustering (K-Means++)", "Unsupervised Partitioning", "src/unsupervised.py", "Latent market tier segmentation (K=4 market clusters)."),
        ("Isolation Forest (iTrees)", "Unsupervised Anomaly", "src/unsupervised.py", "Real-time fraud audit; isolates corrupt/bogus quote submissions.")
    ]
    for i, r in enumerate(rows_data):
        for j, val in enumerate(r):
            cell = table.cell(i + 1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_CARD if (i % 2 == 0) else C_DARK_CARD
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10)
            p.font.color.rgb = C_WHITE if j != 0 else C_EMERALD

    s4.notes_slide.notes_text_frame.text = (
        "[TIMING: 3:15 - 4:15]\n"
        "TALKING SCRIPT:\n"
        "'Rather than using a single black-box model, ReValue AI systematically implements and benchmarks 11 algorithms:\n"
        "- Supervised Boosting: HistGradientBoosting is our Champion Regressor, leveraging 256 discrete integer bins for split acceleration.\n"
        "- Supervised Bagging: Random Forest with 150 parallel trees decorrelating variance.\n"
        "- Regularized Linear: Ridge L2 for multicollinearity and Lasso L1 for coordinate-descent feature sparsity.\n"
        "- Unsupervised: K-Means++ for latent market tiers and Isolation Forest as our live fraud guardrail.\n"
        "- Statistical Inference: Paired Student's t-test and Kolmogorov-Smirnov hypothesis testing.'"
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 5: Supervised Learning — Linear vs Tree Ensembles (CO-2, CO-3)
    # ─────────────────────────────────────────────────────────────
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "CO-2 & CO-3: Supervised Regression Benchmark", "Linear Models vs. Tree Ensembles: The Mathematical Split")

    # Card 1: Linear
    c1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = C_CARD
    c1.line.color.rgb = C_GOLD
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "Linear Models (CO-2)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_GOLD

    pts1 = [
        "• OLS Normal Equations: w = (X^T X)^-1 X^T y. Suffers severe multicollinearity (RAM, Storage, Retail Price).",
        "• Ridge (L2 Penalty): Adds alpha * ||w||_2^2. Guarantees (X^T X + alpha I) is strictly invertible and well-conditioned.",
        "• Lasso (L1 Penalty): Adds alpha * ||w||_1. Sharp diamond corners force non-informative coefficients to exactly 0.",
        "• Mandatory StandardScaler: Isotropic scaling ensures penalization is uniform across gigabytes and binary wear flags.",
        "• Limitation: Unable to model non-linear hardware cliffs (e.g. sub-80% battery threshold)."
    ]
    for pt in pts1:
        p_sub = tf1.add_paragraph()
        p_sub.text = pt
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = C_WHITE
        p_sub.space_before = Pt(8)

    # Card 2: Trees
    c2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = C_CARD
    c2.line.color.rgb = C_EMERALD
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Tree Ensembles (CO-3)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD

    pts2 = [
        "• Random Forest: Bagging ensemble of 150 deep, low-bias trees with sqrt(p) feature subsampling to minimize correlation.",
        "• HistGradientBoosting: Sequential additive boosting with shrinkage (eta = 0.08) fitting pseudo-residuals.",
        "• Histogram Binning: Bins continuous features into 256 integer bins, accelerating node split evaluation by 10x.",
        "• Axis-Aligned Thresholds: Naturally isolates non-linear cliff degradation points without requiring feature polynomial expansions.",
        "• Monotonic Invariance: Split choices depend strictly on rank ordering, rendering trees naturally scale-invariant."
    ]
    for pt in pts2:
        p_sub = tf2.add_paragraph()
        p_sub.text = pt
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = C_WHITE
        p_sub.space_before = Pt(8)

    s5.notes_slide.notes_text_frame.text = (
        "[TIMING: 4:15 - 5:30]\n"
        "TALKING SCRIPT:\n"
        "'Under CO-2 and CO-3, we demonstrate why tree ensembles dominate linear models on tabular electronics data.\n"
        "In linear regression, multicollinearity between RAM, storage, and launch price causes the OLS matrix inversion to become ill-conditioned. "
        "Ridge regularization resolves this by adding a spherical L2 penalty alpha * ||w||_2^2, making the matrix strictly invertible.\n"
        "However, linear hyperplanes cannot capture the sharp cliff where a laptop battery drops below 80% and loses 30% of its value overnight. "
        "HistGradientBoosting solves this by constructing axis-aligned orthogonal splits along the degradation boundary. "
        "By binning features into 256 discrete bins, split finding speeds up by an order of magnitude.'"
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 6: Empirical Benchmark Results & Hypothesis Testing (CO-5)
    # ─────────────────────────────────────────────────────────────
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "CO-5: Model Evaluation, Selection & Verification", "Empirical Benchmark Leaderboard & Statistical Significance")

    # Table
    tbl2_shape = s6.shapes.add_table(7, 5, Inches(0.8), Inches(1.8), Inches(11.7), Inches(3.2))
    tbl2 = tbl2_shape.table
    tbl2.columns[0].width = Inches(3.2)
    tbl2.columns[1].width = Inches(2.0)
    tbl2.columns[2].width = Inches(2.0)
    tbl2.columns[3].width = Inches(2.0)
    tbl2.columns[4].width = Inches(2.5)

    headers2 = ["Model", "Test MAE ($)", "Test MAE (₹)", "Test R² Score", "Performance vs Baseline"]
    for j, h in enumerate(headers2):
        cell = tbl2.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_DARK_CARD
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_GOLD

    b_rows = [
        ("HistGradientBoosting (Champion)", "$27.44", "₹2,305", "0.9877", "66.1% Error Reduction (p < 10^-48)"),
        ("Random Forest", "$41.19", "₹3,460", "0.9738", "49.1% Error Reduction (p < 10^-27)"),
        ("Ridge Regression (L2)", "$80.93", "₹6,798", "0.9048", "Regularized Baseline (alpha = 1.0)"),
        ("Linear Regression (OLS)", "$80.90", "₹6,796", "0.9052", "Unregularized Closed-Form"),
        ("Lasso Regression (L1)", "$82.02", "₹6,890", "0.9034", "Feature Sparsity Baseline"),
        ("ElasticNet", "$117.30", "₹9,853", "0.8060", "Convex Combination (l1-ratio = 0.5)")
    ]
    for i, r in enumerate(b_rows):
        for j, val in enumerate(r):
            cell = tbl2.cell(i + 1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_CARD if (i % 2 == 0) else C_DARK_CARD
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10)
            p.font.color.rgb = C_EMERALD if i == 0 else C_WHITE

    # Callout card for Paired t-test
    ttest_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.6))
    ttest_card.fill.solid()
    ttest_card.fill.fore_color.rgb = C_CARD
    ttest_card.line.color.rgb = C_EMERALD
    tf_t = ttest_card.text_frame
    tf_t.word_wrap = True
    tp1 = tf_t.paragraphs[0]
    tp1.text = "🔬 Scientific Verification: Paired Student's t-Test on Test Errors"
    tp1.font.size = Pt(13)
    tp1.font.bold = True
    tp1.font.color.rgb = C_GOLD

    tp2 = tf_t.add_paragraph()
    tp2.text = (
        "• Null Hypothesis H0: Mean absolute error of HistGBDT equals Ridge baseline (mu_HistGBDT = mu_Ridge).\n"
        "• Empirical Result: t-statistic = 16.82, p-value = 2.8 x 10^-49 << 0.01.\n"
        "• Conclusion: We reject H0 with extreme statistical significance. HistGradientBoosting's $53.49 (₹4,493) average margin per device is scientifically proven superior."
    )
    tp2.font.size = Pt(11)
    tp2.font.color.rgb = C_WHITE
    tp2.space_before = Pt(4)

    s6.notes_slide.notes_text_frame.text = (
        "[TIMING: 5:30 - 6:45]\n"
        "TALKING SCRIPT:\n"
        "'Under CO-5, we evaluated all models on an untouched 15% held-out test set of 375 devices. "
        "HistGradientBoosting achieved a Test MAE of ₹2,305 ($27.44) and an R² score of 0.9877, beating Ridge by 66.1%.\n"
        "Critically, in engineering we do not rely on raw metrics alone. We conducted a Paired Student's t-test comparing sample-by-sample errors. "
        "The resulting p-value is 2.8 x 10^-49, allowing us to reject the null hypothesis and scientifically prove "
        "that our model's superiority is genuine and not an artifact of random sample variance.'"
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 7: Unsupervised Intelligence — K-Means & Fraud Guardrails (CO-4)
    # ─────────────────────────────────────────────────────────────
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "CO-4: Unsupervised Learning & Anomaly Guardrails", "Latent Market Segmentation & Real-Time Fraud Auditing")

    u_cols = [
        ("K-Means++ Clustering", [
            "• Centroid initialization via D(x)^2 probability distribution prevents poor local minima.",
            "• Evaluated across Elbow curve inertia & Silhouette coefficient peaks.",
            "• Identifies K = 4 latent market tiers: Budget Entry, Workhorse, Depreciated Flagship, Luxury.",
            "• Powers automatic tier badges in valuation reports."
        ], C_GOLD),
        ("Principal Component Analysis (PCA)", [
            "• Eigen-decomposition of the covariance matrix of 9 hardware specifications.",
            "• Top 2 principal components capture 67.7% of total dataset variance.",
            "• Enables 2D/3D cluster scatter visualization on Tab 3 of the live platform.",
            "• Confirms continuous separation between depreciated and flagship hardware."
        ], C_WHITE),
        ("Isolation Forest (Fraud Guardrail)", [
            "• Isolates anomalies instead of profiling normal clusters using random axis cuts.",
            "• Extreme anomalies require significantly fewer tree splits (shorter path lengths E[h(x)]).",
            "• Contamination rate = 0.04 screens fraudulent inputs (e.g. ₹10,000 iPhone 16 Pro Max).",
            "• Real-time guardrail executes prior to regressor inference."
        ], C_EMERALD)
    ]
    for i, (uc_title, uc_bullets, uc_color) in enumerate(u_cols):
        u_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 4.0), Inches(1.8), Inches(3.7), Inches(4.8))
        u_box.fill.solid()
        u_box.fill.fore_color.rgb = C_CARD
        u_box.line.color.rgb = uc_color
        utf = u_box.text_frame
        utf.word_wrap = True
        up = utf.paragraphs[0]
        up.text = uc_title
        up.font.size = Pt(15)
        up.font.bold = True
        up.font.color.rgb = uc_color

        for b in uc_bullets:
            ubp = utf.add_paragraph()
            ubp.text = b
            ubp.font.size = Pt(11)
            ubp.font.color.rgb = C_WHITE
            ubp.space_before = Pt(8)

    s7.notes_slide.notes_text_frame.text = (
        "[TIMING: 6:45 - 7:45]\n"
        "TALKING SCRIPT:\n"
        "'Under CO-4, ReValue AI incorporates two unsupervised algorithms:\n"
        "First, K-Means++ clustering partitions devices into 4 latent market tiers: Budget Entry, Daily Workhorse, Depreciated Flagship, and Luxury. "
        "We validated optimal K using Elbow inertia and Silhouette scores, and visualised the clusters via 2D PCA capturing 67.7% variance.\n"
        "Second, Isolation Forest serves as our live fraud guardrail. Unlike clustering, Isolation Forest isolates outliers by constructing random binary trees. "
        "Because fraudulent quotes have extreme attribute combinations, they require far fewer random splits to isolate. "
        "With a contamination rate of 0.04, it intercepts bogus listings before inference touches the pricing engine.'"
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 8: Production Serving & MLOps Drift Telemetry (CO-6)
    # ─────────────────────────────────────────────────────────────
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "CO-6: ML Engineering — From Notebook to Production", "FastAPI Serving, Kolmogorov-Smirnov Drift & Retraining")

    m_items = [
        ("FastAPI Microservice (src/api.py)", [
            "• Endpoints: POST /predict, POST /predict/batch, GET /health, GET /monitor/drift.",
            "• Pydantic v2 data contracts enforce strict type checking and range limits.",
            "• Sub-10ms response latency; outputs 90% confidence intervals via residual standard deviation."
        ]),
        ("Kolmogorov-Smirnov Covariate Drift (src/monitoring.py)", [
            "• Continuous two-sample KS hypothesis test: D = sup_x |F_prod(x) - F_ref(x)|.",
            "• Compares empirical cumulative distribution of live production requests against reference training data.",
            "• Automatically flags covariate shift when p < 0.05 on 2 or more features."
        ]),
        ("Champion-Challenger Retraining Loop (src/retrain.py)", [
            "• Drift alerts trigger automated offline retraining on accumulated SQLite telemetry.",
            "• Candidate challenger models must beat the Champion's held-out Test MAE before automated redeployment."
        ])
    ]
    for i, (m_title, m_bullets) in enumerate(m_items):
        box = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8 + i * 1.65), Inches(11.7), Inches(1.45))
        box.fill.solid()
        box.fill.fore_color.rgb = C_CARD
        box.line.color.rgb = C_EMERALD if i == 1 else (C_GOLD if i == 0 else RGBColor(40, 50, 75))
        mtf = box.text_frame
        mtf.word_wrap = True
        mp = mtf.paragraphs[0]
        mp.text = m_title
        mp.font.size = Pt(14)
        mp.font.bold = True
        mp.font.color.rgb = C_GOLD if i == 0 else C_EMERALD

        for b in m_bullets:
            mbp = mtf.add_paragraph()
            mbp.text = b
            mbp.font.size = Pt(11)
            mbp.font.color.rgb = C_WHITE
            mbp.space_before = Pt(3)

    s8.notes_slide.notes_text_frame.text = (
        "[TIMING: 7:45 - 8:45]\n"
        "TALKING SCRIPT:\n"
        "'Under CO-6, we transition from modeling to production ML engineering.\n"
        "Our FastAPI backend serves predictions with Pydantic payload validation and sub-10ms latency. "
        "Every inference is logged into an SQLite database with timestamp, input vector, and output bounds.\n"
        "To handle data distribution shifts over time, we implemented the non-parametric Two-Sample Kolmogorov-Smirnov test. "
        "It monitors continuous maximum CDF distance. If p < 0.05 on two or more features, the system flags covariate drift "
        "and triggers the Champion-Challenger retraining pipeline in src/retrain.py, redeploying only if the challenger outperforms our current benchmark.'"
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 9: Live Interactive Platform Showcase
    # ─────────────────────────────────────────────────────────────
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Production Deployment & Live Studio", "Interactive Experience & Production URLs")

    # Left: Links & Specs
    p_box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    p_box.fill.solid()
    p_box.fill.fore_color.rgb = C_CARD
    p_box.line.color.rgb = C_EMERALD
    ptf = p_box.text_frame
    ptf.word_wrap = True
    pp = ptf.paragraphs[0]
    pp.text = "🌐 Live Deployment Endpoints"
    pp.font.size = Pt(16)
    pp.font.bold = True
    pp.font.color.rgb = C_EMERALD

    pts_p = [
        "• Live Streamlit Cloud Studio:\n  https://revalueai-xhrrkbmbuexrswr3ermwnf.streamlit.app/",
        "• GitHub Source Repository:\n  https://github.com/dhanush-chanchali/ReValue_AI",
        "• Automated Unit Test Suite:\n  7/7 passing unit & integration tests (pytest tests/ -v)",
        "• FastAPI REST Microservice:\n  FastAPI on port 8000 with Swagger UI at /docs",
        "• Unified One-Command Starter:\n  ./start.sh orchestrates both backend and frontend."
    ]
    for pt in pts_p:
        p_sub = ptf.add_paragraph()
        p_sub.text = pt
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = C_WHITE
        p_sub.space_before = Pt(8)

    # Right: Studio Features
    f_box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8))
    f_box.fill.solid()
    f_box.fill.fore_color.rgb = C_CARD
    f_box.line.color.rgb = C_GOLD
    ftf = f_box.text_frame
    ftf.word_wrap = True
    fp = ftf.paragraphs[0]
    fp.text = "🖥️ In-App Capabilities (5 Tabs)"
    fp.font.size = Pt(16)
    fp.font.bold = True
    fp.font.color.rgb = C_GOLD

    pts_f = [
        "1. 📱 Tab 1 (Sell Smartphone): 3-Screen wizard, release-year guardrails, 60fps battery slider, live ₹ pricing.",
        "2. 💻 Tab 2 (Sell Laptop): Specs configurator, processor tiers, SSD wear flags, confidence intervals.",
        "3. 📊 Tab 3 (ML Model Insights): Benchmarks, residual homoscedasticity, 2D PCA cluster map.",
        "4. 🛡️ Tab 4 (MLOps Command Centre): Real-time SQLite request logs, KS-test telemetry, one-click retrain.",
        "5. 📘 Tab 5 (Syllabus Coverage): KaTeX vector math viewer for all 20 curriculum concepts and 11 algorithms."
    ]
    for pt in pts_f:
        f_sub = ftf.add_paragraph()
        f_sub.text = pt
        f_sub.font.size = Pt(11)
        f_sub.font.color.rgb = C_WHITE
        f_sub.space_before = Pt(6)

    s9.notes_slide.notes_text_frame.text = (
        "[TIMING: 8:45 - 9:30]\n"
        "TALKING SCRIPT:\n"
        "'The platform is live and accessible online at revalueai-xhrrkbmbuexrswr3ermwnf.streamlit.app. "
        "It features five interactive tabs including guided trade-in wizards, deep residual diagnostics, "
        "an MLOps command centre for live drift telemetry, and an academic syllabus explorer rendering textbook-grade KaTeX vector math.\n"
        "All seven unit and integration tests are passing, and our entire service stack boots with a single command: ./start.sh.'"
    )

    # ─────────────────────────────────────────────────────────────
    # SLIDE 10: Conclusion, Viva Defense & Q&A
    # ─────────────────────────────────────────────────────────────
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Summary & Conclusion", "Key Takeaways & Academic Contribution")

    c_box = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = C_CARD
    c_box.line.color.rgb = C_GOLD
    ctf = c_box.text_frame
    ctf.word_wrap = True
    cp = ctf.paragraphs[0]
    cp.text = "🎯 Summary of Contributions"
    cp.font.size = Pt(18)
    cp.font.bold = True
    cp.font.color.rgb = C_GOLD

    summary_pts = [
        "1. Solved Market Asymmetry: Built an empirical, fair-market electronics resale valuation engine trained across 2,500+ device profiles.",
        "2. Proven Model Superiority: Champion HistGradientBoosting achieves R² = 0.9877 and ₹2,305 MAE, beating linear baselines by 66.1% (p < 10^-48).",
        "3. Multi-Layer Guardrails: Integrated Isolation Forest (contamination = 0.04) and K-Means++ clustering for real-time fraud and market tier segmentation.",
        "4. Production Readiness: FastAPI microservice, Streamlit Cloud live deployment, and Kolmogorov-Smirnov covariate drift telemetry.",
        "5. Full Curriculum Coverage: Seamless execution across modules CO-1 through CO-6 with zero training-serving skew."
    ]
    for pt in summary_pts:
        p_sub = ctf.add_paragraph()
        p_sub.text = pt
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = C_WHITE
        p_sub.space_before = Pt(10)

    p_end = ctf.add_paragraph()
    p_end.text = "Thank you! I welcome your questions and viva evaluation."
    p_end.font.size = Pt(14)
    p_end.font.bold = True
    p_end.font.color.rgb = C_EMERALD
    p_end.space_before = Pt(16)
    p_end.alignment = PP_ALIGN.CENTER

    s10.notes_slide.notes_text_frame.text = (
        "[TIMING: 9:30 - 10:00]\n"
        "TALKING SCRIPT:\n"
        "'In conclusion, ReValue AI bridges the gap between academic Machine Learning theory and production software engineering. "
        "It achieves high empirical accuracy with R² = 0.9877, protects the platform with unsupervised fraud guardrails, "
        "and implements proactive MLOps drift telemetry.\n"
        "Thank you for your time and attention. I am now open to your questions and viva defense.'"
    )

    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "ReValue_AI_Presentation.pptx"
    create_deck(out)
